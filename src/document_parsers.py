import os
import re
import logging
from typing import List, Dict, Optional
from src.interfaces import DocumentParser, DocumentFormat, ParsedDocument

logger = logging.getLogger(__name__)


class TxtParser(DocumentParser):
    def supported_formats(self) -> List[DocumentFormat]:
        return [DocumentFormat.TXT]

    def can_parse(self, source: str) -> bool:
        return source.lower().endswith(".txt")

    def parse(self, source: str, **kwargs) -> ParsedDocument:
        encodings = kwargs.get("encodings", ["utf-8", "gbk", "gb2312", "utf-16"])
        raw_text = ""
        for enc in encodings:
            try:
                with open(source, "r", encoding=enc) as f:
                    raw_text = f.read()
                if raw_text:
                    break
            except (UnicodeDecodeError, UnicodeError):
                continue

        title = os.path.splitext(os.path.basename(source))[0]
        sections = self._split_sections(raw_text)

        return ParsedDocument(
            source_path=source,
            source_format=DocumentFormat.TXT,
            title=title,
            raw_text=raw_text,
            sections=sections,
        )

    def _split_sections(self, text: str) -> List[Dict]:
        sections = []
        current_chapter = ""
        current_article_num = None
        buffer = []

        for line in text.split("\n"):
            stripped = line.strip()
            if not stripped:
                continue

            chapter_match = re.match(r"^第[一二三四五六七八九十百]+章\s+(.+)$", stripped)
            if chapter_match:
                if current_article_num and buffer:
                    sections.append({
                        "chapter": current_chapter,
                        "article_number": current_article_num,
                        "article_text": "".join(buffer).strip(),
                    })
                current_chapter = stripped
                current_article_num = None
                buffer = []
                continue

            article_match = re.match(r"^第([一二三四五六七八九十百]+)条\s*(.*)$", stripped)
            if article_match:
                if current_article_num and buffer:
                    sections.append({
                        "chapter": current_chapter,
                        "article_number": current_article_num,
                        "article_text": "".join(buffer).strip(),
                    })
                current_article_num = article_match.group(1)
                buffer = [article_match.group(2)]
                continue

            if current_article_num:
                buffer.append(stripped)

        if current_article_num and buffer:
            sections.append({
                "chapter": current_chapter,
                "article_number": current_article_num,
                "article_text": "".join(buffer).strip(),
            })

        return sections


class MdParser(DocumentParser):
    def supported_formats(self) -> List[DocumentFormat]:
        return [DocumentFormat.MD]

    def can_parse(self, source: str) -> bool:
        return source.lower().endswith(".md")

    def parse(self, source: str, **kwargs) -> ParsedDocument:
        with open(source, "r", encoding="utf-8") as f:
            raw_text = f.read()

        title = os.path.splitext(os.path.basename(source))[0]
        clean_text = self._strip_markdown(raw_text)
        sections = TxtParser()._split_sections(clean_text)

        tables = self._extract_tables(raw_text)

        return ParsedDocument(
            source_path=source,
            source_format=DocumentFormat.MD,
            title=title,
            raw_text=clean_text,
            sections=sections,
            tables=tables,
        )

    def _strip_markdown(self, text: str) -> str:
        text = re.sub(r"!\[.*?\]\(.*?\)", "", text)
        text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
        text = re.sub(r"^[#]+\s+", "", text, flags=re.MULTILINE)
        text = re.sub(r"\*{1,2}([^*]+)\*{1,2}", r"\1", text)
        text = re.sub(r"`{1,3}[^`]*`{1,3}", "", text)
        return text.strip()

    def _extract_tables(self, text: str) -> List[List[List[str]]]:
        tables = []
        table_pattern = re.compile(r"(\|.+\|\n\|[-:\s|]+\|\n(?:\|.+\|\n?)+)", re.MULTILINE)
        for match in table_pattern.finditer(text):
            rows = []
            for line in match.group(1).strip().split("\n"):
                if re.match(r"^\|[-:\s|]+\|$", line.strip()):
                    continue
                cells = [c.strip() for c in line.strip().split("|")[1:-1]]
                if cells:
                    rows.append(cells)
            if rows:
                tables.append(rows)
        return tables


class PdfParser(DocumentParser):
    def supported_formats(self) -> List[DocumentFormat]:
        return [DocumentFormat.PDF]

    def can_parse(self, source: str) -> bool:
        return source.lower().endswith(".pdf")

    def parse(self, source: str, **kwargs) -> ParsedDocument:
        raw_text = self._extract_text(source)
        images = self._extract_images_info(source)
        tables = self._extract_tables_from_pdf(source)

        title = os.path.splitext(os.path.basename(source))[0]
        sections = TxtParser()._split_sections(raw_text) if raw_text else []

        return ParsedDocument(
            source_path=source,
            source_format=DocumentFormat.PDF,
            title=title,
            raw_text=raw_text,
            sections=sections,
            tables=tables,
            images=images,
            metadata={"page_count": self._get_page_count(source)},
        )

    def _extract_text(self, source: str) -> str:
        try:
            import pymupdf
            doc = pymupdf.open(source)
            text_parts = []
            for page in doc:
                text_parts.append(page.get_text())
            doc.close()
            return "\n\n".join(text_parts)
        except ImportError:
            try:
                import fitz
                doc = fitz.open(source)
                text_parts = []
                for page in doc:
                    text_parts.append(page.get_text())
                doc.close()
                return "\n\n".join(text_parts)
            except ImportError:
                logger.warning("pymupdf/fitz 未安装，PDF文本提取不可用，尝试pdfminer")
                return self._extract_with_pdfminer(source)

    def _extract_with_pdfminer(self, source: str) -> str:
        try:
            from pdfminer.high_level import extract_text
            return extract_text(source)
        except ImportError:
            logger.error("无可用PDF解析库(pymupdf/pdfminer)，请安装其一")
            return ""

    def _extract_images_info(self, source: str) -> List[Dict]:
        images = []
        try:
            import pymupdf
            doc = pymupdf.open(source)
            for page_num, page in enumerate(doc):
                image_list = page.get_images(full=True)
                for img_idx, img in enumerate(image_list):
                    xref = img[0]
                    base_image = doc.extract_image(xref)
                    if base_image:
                        images.append({
                            "page": page_num + 1,
                            "index": img_idx,
                            "width": base_image.get("width", 0),
                            "height": base_image.get("height", 0),
                            "format": base_image.get("ext", ""),
                            "size_bytes": len(base_image.get("image", b"")),
                            "ocr_text": "",
                        })
            doc.close()
        except ImportError:
            pass
        return images

    def _extract_tables_from_pdf(self, source: str) -> List[List[List[str]]]:
        tables = []
        try:
            import pymupdf
            doc = pymupdf.open(source)
            for page in doc:
                tabs = page.find_tables()
                for tab in tabs:
                    rows = []
                    for row in tab.extract():
                        cells = [str(c) if c else "" for c in row]
                        rows.append(cells)
                    if rows:
                        tables.append(rows)
            doc.close()
        except (ImportError, Exception) as e:
            logger.debug(f"PDF表格提取跳过: {e}")
        return tables

    def _get_page_count(self, source: str) -> int:
        try:
            import pymupdf
            doc = pymupdf.open(source)
            count = len(doc)
            doc.close()
            return count
        except ImportError:
            return 0


class DocxParser(DocumentParser):
    def supported_formats(self) -> List[DocumentFormat]:
        return [DocumentFormat.DOCX]

    def can_parse(self, source: str) -> bool:
        return source.lower().endswith(".docx")

    def parse(self, source: str, **kwargs) -> ParsedDocument:
        raw_text = ""
        tables = []
        images = []

        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(source)

            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text.strip())
            raw_text = "\n\n".join(paragraphs)

            for table in doc.tables:
                rows = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    rows.append(cells)
                if rows:
                    tables.append(rows)

            for rel in doc.part.rels.values():
                if "image" in rel.reltype:
                    images.append({
                        "type": "embedded",
                        "format": rel.target_ref.split(".")[-1] if "." in rel.target_ref else "",
                    })
        except ImportError:
            logger.error("python-docx 未安装，Word文档解析不可用")
        except Exception as e:
            logger.error(f"Word文档解析失败: {e}")

        title = os.path.splitext(os.path.basename(source))[0]
        sections = TxtParser()._split_sections(raw_text) if raw_text else []

        return ParsedDocument(
            source_path=source,
            source_format=DocumentFormat.DOCX,
            title=title,
            raw_text=raw_text,
            sections=sections,
            tables=tables,
            images=images,
        )


class DocParser(DocumentParser):
    def supported_formats(self) -> List[DocumentFormat]:
        return [DocumentFormat.DOC]

    def can_parse(self, source: str) -> bool:
        return source.lower().endswith(".doc") and not source.lower().endswith(".docx")

    def parse(self, source: str, **kwargs) -> ParsedDocument:
        docx_path = self._convert_to_docx(source)
        if docx_path:
            docx_parser = DocxParser()
            result = docx_parser.parse(docx_path, **kwargs)
            result.source_format = DocumentFormat.DOC
            result.source_path = source
            try:
                os.remove(docx_path)
            except:
                pass
            return result
        else:
            logger.warning(f"无法转换.doc文件，尝试直接读取: {source}")
            raw_text = self._fallback_read(source)
            title = os.path.splitext(os.path.basename(source))[0]
            sections = TxtParser()._split_sections(raw_text) if raw_text else []
            return ParsedDocument(
                source_path=source,
                source_format=DocumentFormat.DOC,
                title=title,
                raw_text=raw_text,
                sections=sections,
            )

    def _convert_to_docx(self, source: str) -> Optional[str]:
        import subprocess
        import tempfile
        output_dir = tempfile.mkdtemp()
        try:
            result = subprocess.run(
                ["libreoffice", "--headless", "--convert-to", "docx", "--outdir", output_dir, source],
                capture_output=True, text=True, timeout=60
            )
            if result.returncode == 0:
                basename = os.path.splitext(os.path.basename(source))[0]
                docx_path = os.path.join(output_dir, f"{basename}.docx")
                if os.path.exists(docx_path):
                    return docx_path
        except FileNotFoundError:
            logger.warning("LibreOffice未安装，尝试antiword")
            return self._convert_with_antiword(source, output_dir)
        except subprocess.TimeoutExpired:
            logger.warning("LibreOffice转换超时")
        except Exception as e:
            logger.warning(f"LibreOffice转换失败: {e}")
        return None

    def _convert_with_antiword(self, source: str, output_dir: str) -> Optional[str]:
        import subprocess
        try:
            result = subprocess.run(
                ["antiword", source],
                capture_output=True, text=True, timeout=30
            )
            if result.returncode == 0 and result.stdout.strip():
                txt_path = os.path.join(output_dir, "converted.txt")
                with open(txt_path, "w", encoding="utf-8") as f:
                    f.write(result.stdout)
                return txt_path
        except FileNotFoundError:
            logger.warning("antiword未安装")
        except Exception as e:
            logger.warning(f"antiword转换失败: {e}")
        return None

    def _fallback_read(self, source: str) -> str:
        try:
            with open(source, "rb") as f:
                content = f.read()
            text = content.decode("utf-8", errors="ignore")
            text = re.sub(r'[^\u4e00-\u9fff\u3000-\u303f\uff00-\uffefa-zA-Z0-9\s。，、；：？！""''（）【】《》—…·\n\r-]', '', text)
            return text.strip()
        except Exception as e:
            logger.error(f"无法读取.doc文件: {e}")
            return ""


class ImageParser(DocumentParser):
    def supported_formats(self) -> List[DocumentFormat]:
        return [DocumentFormat.IMAGE]

    def can_parse(self, source: str) -> bool:
        return source.lower().endswith((".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp"))

    def parse(self, source: str, **kwargs) -> ParsedDocument:
        ocr_text = self._ocr(source)
        title = os.path.splitext(os.path.basename(source))[0]

        return ParsedDocument(
            source_path=source,
            source_format=DocumentFormat.IMAGE,
            title=title,
            raw_text=ocr_text,
            sections=[{"chapter": "", "article_number": "img_1", "article_text": ocr_text}] if ocr_text else [],
            images=[{
                "path": source,
                "ocr_text": ocr_text,
                "format": os.path.splitext(source)[1].lstrip("."),
            }],
        )

    def _ocr(self, source: str) -> str:
        try:
            import dashscope
            from dashscope import MultiModalConversation
            dashscope.api_key = __import__("config").DASHSCOPE_API_KEY

            if not dashscope.api_key:
                return ""

            messages = [{
                "role": "user",
                "content": [
                    {"image": f"file://{source}"},
                    {"text": "请提取图片中的所有文字内容，按原文格式输出。如果是法规文档，请保留条文编号。"},
                ],
            }]
            resp = MultiModalConversation.call(
                model="qwen-vl-plus",
                messages=messages,
            )
            if resp.status_code == 200:
                return resp.output.choices[0].message.content
        except ImportError:
            logger.debug("dashscope 多模态不可用")
        except Exception as e:
            logger.debug(f"图片OCR失败: {e}")

        try:
            import pytesseract
            from PIL import Image
            img = Image.open(source)
            return pytesseract.image_to_string(img, lang="chi_sim+eng")
        except ImportError:
            pass
        except Exception as e:
            logger.debug(f"Tesseract OCR失败: {e}")

        return ""


class DocumentParserRegistry:
    def __init__(self):
        self._parsers: Dict[DocumentFormat, DocumentParser] = {}
        self._register_defaults()

    def _register_defaults(self):
        self.register(TxtParser())
        self.register(MdParser())
        self.register(PdfParser())
        self.register(DocxParser())
        self.register(DocParser())
        self.register(ImageParser())

    def register(self, parser: DocumentParser):
        for fmt in parser.supported_formats():
            self._parsers[fmt] = parser

    def get_parser(self, source: str) -> Optional[DocumentParser]:
        for fmt, parser in self._parsers.items():
            if parser.can_parse(source):
                return parser
        return None

    def parse(self, source: str, **kwargs) -> ParsedDocument:
        parser = self.get_parser(source)
        if not parser:
            raise ValueError(f"不支持的文档格式: {source}")
        return parser.parse(source, **kwargs)

    def supported_formats(self) -> List[str]:
        return [fmt.value for fmt in self._parsers.keys()]


parser_registry = DocumentParserRegistry()
